"""Generate one consolidated Excel report per milestone; columns adapt to criteria across projects."""
import re
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE_COLUMNS = ["Participant Name", "Submission File", "Project", "Milestone"]
SCORE_COLUMNS = ["Total Score", "Maximum Score"]
FEEDBACK_COLUMNS = ["Area of Strength", "Area for Improvement", "Overall Feedback"]

def _participant_name(source_file: str) -> str:
    path = Path(source_file)
    match = re.search(r"(Participant(?:[_ -]M[1-4])?[_ -]*\d+)", path.stem, re.IGNORECASE)
    if not match:
        match = re.search(r"(Participant[_ -]*\d+)", path.as_posix(), re.IGNORECASE)
    if match:
        return re.sub(r"[ -]+", "_", match.group(1)).replace("__", "_")
    return path.parent.name or path.stem

def _feedback_text(value):
    if value is None:
        return ""
    if isinstance(value, (list, tuple)):
        return "\n".join(str(item) for item in value if item is not None)
    return str(value)

def _format_worksheet(ws):
    ws.freeze_panes = "A2"
    if ws.max_row >= 1 and ws.max_column >= 1:
        ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for column_cells in ws.columns:
        letter = get_column_letter(column_cells[0].column)
        header = str(column_cells[0].value or "")
        max_len = max((len(str(c.value or "")) for c in column_cells), default=12)
        if header in FEEDBACK_COLUMNS:
            width = 48
        elif header == "Submission File":
            width = 48
        else:
            width = min(38, max(12, max_len + 2))
        ws.column_dimensions[letter].width = width
        for cell in column_cells:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

def _milestone_row(result, source_file):
    milestone = str(result.get("milestone_id", "Milestone_1"))
    row = {
        "Participant Name": _participant_name(source_file),
        "Submission File": str(source_file).replace("\\", "/"),
        "Project": result.get("project_id", ""),
        "Milestone": milestone,
        "Total Score": result.get("total_marks", 0),
        "Maximum Score": result.get("maximum_marks", 100),
        "Area of Strength": _feedback_text(result.get("strengths", [])),
        "Area for Improvement": _feedback_text(result.get("areas_for_improvement", [])),
        "Overall Feedback": _feedback_text(result.get("overall_feedback", "")),
    }
    for criterion in result.get("criteria", []):
        name = str(criterion.get("name", "Criterion")).strip()
        maximum = criterion.get("max_marks", "")
        row[f"{name} ({maximum})"] = criterion.get("marks_awarded", "")
    return row

def write_reports(result, source_file, output_dir):
    """Upsert submission into one milestone workbook; preserve union of every project's rubric columns."""
    out = Path(output_dir)
    milestone = str(result.get("milestone_id", "Milestone_1"))
    match = re.search(r"([1-4])$", milestone)
    number = match.group(1) if match else "1"
    report_folder = out / f"Submission-Milestone{number}-Report"
    report_folder.mkdir(parents=True, exist_ok=True)
    report_path = report_folder / f"Submission-Milestone{number}-Report.xlsx"

    new_row = _milestone_row(result, source_file)
    old_rows = []
    old_columns = []
    if report_path.exists():
        try:
            existing = pd.read_excel(report_path, sheet_name="Feedback Report", dtype=object).fillna("")
            old_columns = list(existing.columns)
            old_rows = existing.to_dict(orient="records")
        except (ValueError, OSError, ImportError):
            old_rows, old_columns = [], []

    key = (str(new_row["Submission File"]), str(new_row["Milestone"]))
    old_rows = [
        row for row in old_rows
        if (str(row.get("Submission File", "")), str(row.get("Milestone", ""))) != key
    ]
    old_rows.append(new_row)

    standard = BASE_COLUMNS + SCORE_COLUMNS + FEEDBACK_COLUMNS
    criterion_columns = [c for c in old_columns if c not in standard]
    criterion_columns.extend(
        c for c in new_row
        if c not in standard and c not in criterion_columns
    )
    columns = BASE_COLUMNS + criterion_columns + SCORE_COLUMNS + FEEDBACK_COLUMNS
    frame = pd.DataFrame(old_rows).reindex(columns=columns).fillna("")
    if not frame.empty:
        frame = frame.drop_duplicates(subset=["Submission File", "Milestone"], keep="last")
        frame = frame.sort_values(["Participant Name", "Submission File"], kind="stable")

    with pd.ExcelWriter(report_path, engine="openpyxl") as writer:
        frame.to_excel(writer, index=False, sheet_name="Feedback Report")
        _format_worksheet(writer.book["Feedback Report"])
    return report_path, None
